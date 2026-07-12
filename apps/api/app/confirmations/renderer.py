from io import BytesIO

from docx import Document as WordDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas


def render_docx(snapshot: dict) -> bytes:
    document = WordDocument()
    document.add_heading("Business Registration Confirmation", level=1)
    company = snapshot["company"]
    document.add_paragraph(f"English name: {company['name_en']}")
    document.add_paragraph(f"Chinese name: {company['name_zh']}")
    document.add_paragraph(f"Business scope: {company['business_scope']}")
    document.add_paragraph(f"Registered capital: {company['registered_capital']} {company['currency']}")
    document.add_heading("Directors", level=2)
    for director in snapshot["directors"]:
        document.add_paragraph(director["name"])
    document.add_heading("Shareholders", level=2)
    for shareholder in snapshot["shareholders"]:
        document.add_paragraph(f"{shareholder['name']}: {shareholder['share_percentage']}%")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(snapshot: dict) -> bytes:
    output = BytesIO()
    canvas = Canvas(output, pagesize=A4)
    company = snapshot["company"]
    lines = [
        "Business Registration Confirmation",
        f"English name: {company['name_en']}",
        f"Chinese name: {company['name_zh']}",
        f"Business scope: {company['business_scope']}",
        f"Registered capital: {company['registered_capital']} {company['currency']}",
    ]
    y = 800
    for line in lines:
        canvas.drawString(60, y, line)
        y -= 28
    canvas.save()
    return output.getvalue()

